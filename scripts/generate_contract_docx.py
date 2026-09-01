from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
CONTENT_PATH = ROOT / "src" / "data" / "contract-template.json"
OUTPUT_PATH = ROOT / "output" / "docx" / "contrato-general-asesoria-academica-tesis20.docx"
PUBLIC_PATH = ROOT / "public" / "downloads" / "contrato-general-asesoria-academica-tesis20.docx"

YELLOW = "F5BD42"
INK = "1E2022"
MUTED = "5F6264"
LIGHT = "F5F3EE"
BORDER = "D8D5CE"
WHITE = "FFFFFF"

YELLOW_RGB = RGBColor(0xF5, 0xBD, 0x42)
INK_RGB = RGBColor(0x1E, 0x20, 0x22)
MUTED_RGB = RGBColor(0x5F, 0x62, 0x64)

BLANK = "____________________"

EXTRA_FILL_FIELDS = [
    "Nombre del alumno: ____________________",
    "Servicio contratado: ____________________",
    "Monto total (S/): ____________________",
    "Fechas (inicio y término): ____________________",
]


def set_run_font(run, *, size=10, bold=False, color=INK_RGB, name="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=240, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Twips(line)
    if align is not None:
        paragraph.alignment = align


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def set_cell_margins(cell, *, top=60, bottom=60, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_fixed(table, width_cm):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(width_cm * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)


def add_bottom_border(paragraph, color=INK, size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_fill_in_run(paragraph, placeholder=""):
    """Underline an empty span so the party can type over it in Word."""
    run = paragraph.add_run(placeholder or BLANK)
    set_run_font(run, size=10, color=INK_RGB)
    run.underline = True
    rpr = run._element.get_or_add_rPr()
    u = rpr.find(qn("w:u"))
    if u is None:
        u = OxmlElement("w:u")
        rpr.append(u)
    u.set(qn("w:val"), "single")
    u.set(qn("w:color"), INK)


def split_labeled_value(raw):
    if ": " not in raw:
        return raw, ""
    label, value = raw.split(": ", 1)
    return label.strip(), value.strip()


def is_blank_value(value):
    stripped = value.replace("_", "").replace(" ", "")
    return stripped == ""


def write_cell_text(cell, text, *, size=9, bold=False, color=INK_RGB, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph_spacing(paragraph, before=0, after=0, line=220, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_labeled_fill_cell(cell, label, value=""):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    shade_cell(cell, LIGHT)
    set_cell_borders(cell)

    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph_spacing(paragraph, before=0, after=0, line=230)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=9, bold=True, color=MUTED_RGB)

    if value and not is_blank_value(value):
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10, color=INK_RGB)
    else:
        add_fill_in_run(paragraph)


def add_header_bar(section):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    set_paragraph_spacing(paragraph, before=0, after=0, line=200)

    table = header.add_table(rows=1, cols=2, width=Cm(18.0))
    set_table_fixed(table, 18.0)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(6.0)
    table.columns[1].width = Cm(12.0)
    prevent_row_split(table.rows[0])

    left, right = table.rows[0].cells
    for cell in (left, right):
        shade_cell(cell, YELLOW)
        set_cell_borders(cell, YELLOW, "0")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    left_p = write_cell_text(left, "TESIS20", size=12, bold=True, color=INK_RGB)
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    right_p = write_cell_text(
        right,
        "CONTRATO GENERAL DE ASESORÍA ACADÉMICA",
        size=8,
        bold=True,
        color=INK_RGB,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    spacer = header.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=0, line=80)


def add_footer_bar(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    set_paragraph_spacing(paragraph, before=4, after=0, line=200)
    add_bottom_border(paragraph, YELLOW, "16")

    brand = paragraph.add_run("TESIS20  ·  Modelo general para Perú — completar antes de firmar")
    set_run_font(brand, size=8, bold=False, color=MUTED_RGB)

    page_run = paragraph.add_run("          Página ")
    set_run_font(page_run, size=8, color=MUTED_RGB)
    add_page_field(paragraph)
    of_run = paragraph.add_run(" de ")
    set_run_font(of_run, size=8, color=MUTED_RGB)
    add_num_pages_field(paragraph)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8, color=MUTED_RGB)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_num_pages_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8, color=MUTED_RGB)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_heading_block(document, data):
    title = document.add_paragraph()
    set_paragraph_spacing(title, before=6, after=4, line=320, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = title.add_run(data["title"].upper())
    set_run_font(run, size=15, bold=True, color=INK_RGB)

    version = document.add_paragraph()
    set_paragraph_spacing(version, before=0, after=10, line=220, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = version.add_run(data["version"])
    set_run_font(run, size=9, bold=True, color=MUTED_RGB)

    intro = document.add_paragraph()
    set_paragraph_spacing(intro, before=0, after=12, line=240)
    run = intro.add_run(data["intro"])
    set_run_font(run, size=10.5, color=INK_RGB)


def add_section_label(document, text):
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=4, after=6, line=240)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=10, bold=True, color=INK_RGB)


def add_fields_table(document, fields):
    add_section_label(document, "Datos a completar")
    hint = document.add_paragraph()
    set_paragraph_spacing(hint, before=0, after=8, line=220)
    run = hint.add_run(
        "Completa en Word los espacios subrayados: nombre del alumno, DNI, "
        "universidad, servicio, monto, fechas y firmas."
    )
    set_run_font(run, size=9, color=MUTED_RGB)

    rows = []
    for index in range(0, len(fields), 2):
        left = fields[index]
        right = fields[index + 1] if index + 1 < len(fields) else ""
        rows.append((left, right))

    table = document.add_table(rows=len(rows), cols=2)
    set_table_fixed(table, 18.0)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(9.0)
    table.columns[1].width = Cm(9.0)

    for row_index, (left, right) in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        for cell, raw in zip(row.cells, (left, right)):
            if not raw:
                shade_cell(cell, LIGHT)
                set_cell_borders(cell)
                write_cell_text(cell, "")
                continue
            label, value = split_labeled_value(raw)
            add_labeled_fill_cell(cell, label, value)

    document.add_paragraph().paragraph_format.space_after = Pt(8)


def add_clauses(document, clauses):
    add_section_label(document, "Cláusulas")
    for clause in clauses:
        title = document.add_paragraph()
        set_paragraph_spacing(title, before=8, after=3, line=240)
        run = title.add_run(clause["title"])
        set_run_font(run, size=11, bold=True, color=INK_RGB)

        for text in clause["paragraphs"]:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph, before=0, after=6, line=240)
            add_text_with_blanks(paragraph, text)


def add_text_with_blanks(paragraph, text):
    """Keep clause blanks underlined so they stay editable in Word."""
    pieces = []
    buffer = ""
    index = 0
    while index < len(text):
        if text[index] == "_":
            if buffer:
                pieces.append(("text", buffer))
                buffer = ""
            underscores = 0
            while index < len(text) and text[index] == "_":
                underscores += 1
                index += 1
            pieces.append(("blank", "_" * max(underscores, 12)))
        else:
            buffer += text[index]
            index += 1
    if buffer:
        pieces.append(("text", buffer))

    if not pieces:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK_RGB)
        return

    for kind, value in pieces:
        if kind == "blank":
            add_fill_in_run(paragraph, value)
        else:
            run = paragraph.add_run(value)
            set_run_font(run, size=10.5, color=INK_RGB)


def add_signatures(document, signatures):
    add_section_label(document, "Firmas")
    hint = document.add_paragraph()
    set_paragraph_spacing(hint, before=0, after=8, line=220)
    run = hint.add_run("Espacios editables para nombre, documento de identidad y firma de cada parte.")
    set_run_font(run, size=9, color=MUTED_RGB)

    table = document.add_table(rows=1, cols=2)
    set_table_fixed(table, 18.0)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(9.0)
    table.columns[1].width = Cm(9.0)
    prevent_row_split(table.rows[0])

    for cell, raw in zip(table.rows[0].cells, signatures):
        shade_cell(cell, WHITE)
        set_cell_borders(cell, BORDER, "8")
        set_cell_margins(cell, top=120, bottom=140, left=120, right=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        lines = raw.split("\n")
        first = cell.paragraphs[0]
        first.clear()
        set_paragraph_spacing(first, before=0, after=10, line=240)
        heading = first.add_run(lines[0])
        set_run_font(heading, size=11, bold=True, color=INK_RGB)

        for line in lines[1:]:
            paragraph = cell.add_paragraph()
            set_paragraph_spacing(paragraph, before=6, after=8, line=260)
            label, value = split_labeled_value(line)
            label_run = paragraph.add_run(f"{label}: ")
            set_run_font(label_run, size=10, bold=True, color=MUTED_RGB)
            if value and not is_blank_value(value):
                value_run = paragraph.add_run(value)
                set_run_font(value_run, size=10, color=INK_RGB)
            else:
                add_fill_in_run(paragraph, BLANK)
            add_bottom_border(paragraph, BORDER, "8")


def add_notice(document, data):
    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, before=10, after=0, line=160)

    legal_labels = "; ".join(source["label"] for source in data["legalSources"])
    legal = document.add_paragraph()
    set_paragraph_spacing(legal, before=8, after=6, line=220)
    bold = legal.add_run("Marco legal referencial: ")
    set_run_font(bold, size=9, bold=True, color=INK_RGB)
    rest = legal.add_run(f"{legal_labels}.")
    set_run_font(rest, size=9, color=MUTED_RGB)

    table = document.add_table(rows=1, cols=1)
    set_table_fixed(table, 18.0)
    cell = table.rows[0].cells[0]
    shade_cell(cell, LIGHT)
    set_cell_borders(cell, YELLOW, "12")
    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph_spacing(paragraph, before=0, after=0, line=230)
    label = paragraph.add_run("Importante: ")
    set_run_font(label, size=9.5, bold=True, color=INK_RGB)
    notice = paragraph.add_run(data["notice"])
    set_run_font(notice, size=9.5, color=INK_RGB)


def collect_fields(data):
    seen_labels = {split_labeled_value(field)[0].casefold() for field in data["fields"]}
    fields = list(data["fields"])
    for extra in EXTRA_FILL_FIELDS:
        label, _ = split_labeled_value(extra)
        if label.casefold() not in seen_labels:
            fields.append(extra)
            seen_labels.add(label.casefold())
    return fields


def build_docx():
    data = json.loads(CONTENT_PATH.read_text(encoding="utf-8"))
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    PUBLIC_PATH.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.core_properties.author = "Tesis20"
    document.core_properties.title = data["title"]
    document.core_properties.subject = "Modelo general de contrato de asesoría académica e investigación"
    document.core_properties.category = "Contrato"
    document.core_properties.language = "es-PE"

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.6)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = INK_RGB

    add_header_bar(section)
    add_footer_bar(section)
    add_heading_block(document, data)
    add_fields_table(document, collect_fields(data))
    add_clauses(document, data["clauses"])
    add_signatures(document, data["signatures"])
    add_notice(document, data)

    document.save(OUTPUT_PATH)
    shutil.copy2(OUTPUT_PATH, PUBLIC_PATH)
    print(f"Generated: {OUTPUT_PATH}")
    print(f"Published: {PUBLIC_PATH}")


if __name__ == "__main__":
    build_docx()
